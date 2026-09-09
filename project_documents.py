"""Gestione locale di immagini, Word e PDF di Scrittore Site.

Le funzioni qui non accedono a Streamlit, al database o ai crediti. In caso di
errore sollevano un'eccezione: l'interfaccia decide come mostrarla all'utente.
"""

from __future__ import annotations

from collections.abc import Callable
from io import BytesIO
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, ns
from docx.shared import Inches, Pt
from PIL import Image


def normalizza_immagine_caricata(file_caricato) -> bytes:
    """Converte un'immagine caricata in PNG RGB compatto e sicuro."""
    sorgente = Image.open(BytesIO(file_caricato.getvalue()))
    if sorgente.mode in ("RGBA", "LA"):
        sfondo = Image.new("RGB", sorgente.size, "white")
        sfondo.paste(sorgente, mask=sorgente.getchannel("A"))
        sorgente = sfondo
    else:
        sorgente = sorgente.convert("RGB")
    sorgente.thumbnail((1400, 1400), Image.Resampling.LANCZOS)
    output = BytesIO()
    sorgente.save(output, format="PNG", optimize=True)
    return output.getvalue()


def elimina_paragrafo_docx(paragrafo) -> None:
    """Elimina un paragrafo vuoto da un documento Word."""
    elemento = paragrafo._element
    elemento.getparent().remove(elemento)
    paragrafo._p = paragrafo._element = None


def aggiungi_numeri_pagina_docx(documento) -> None:
    """Inserisce il campo numero pagina nel piè di pagina Word."""
    for sezione in documento.sections:
        paragrafo = sezione.footer.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        campo_inizio = OxmlElement("w:fldChar")
        campo_inizio.set(ns.qn("w:fldCharType"), "begin")
        istruzione = OxmlElement("w:instrText")
        istruzione.text = "PAGE"
        campo_fine = OxmlElement("w:fldChar")
        campo_fine.set(ns.qn("w:fldCharType"), "end")
        run = paragrafo.add_run()
        run._r.append(campo_inizio)
        run._r.append(istruzione)
        run._r.append(campo_fine)


def formatta_manoscritto_kdp(file_docx, pulisci_testo: Callable[[str], str]) -> BytesIO:
    """Applica il formato Word 6×9 al manoscritto caricato."""
    documento = Document(BytesIO(file_docx.getvalue()))
    for nome_stile in ("Heading 1", "Heading 2"):
        try:
            documento.styles[nome_stile]
        except KeyError:
            documento.styles.add_style(nome_stile, WD_STYLE_TYPE.PARAGRAPH)

    for sezione in documento.sections:
        sezione.page_width = Inches(6)
        sezione.page_height = Inches(9)
        sezione.top_margin = Inches(0.75)
        sezione.bottom_margin = Inches(0.75)
        sezione.left_margin = Inches(0.75)
        sezione.right_margin = Inches(0.75)

    for paragrafo in list(documento.paragraphs):
        testo = pulisci_testo(paragrafo.text).strip()
        if not testo:
            elimina_paragrafo_docx(paragrafo)
            continue
        paragrafo.text = " ".join(testo.split())
        if len(paragrafo.text) < 80 and re.search(r"(?i)\b(capitolo|chapter|parte|part)\b", paragrafo.text):
            paragrafo.style = "Heading 1"
            paragrafo.paragraph_format.page_break_before = True
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.paragraph_format.space_before = Pt(0)
            paragrafo.paragraph_format.space_after = Pt(30)
        elif len(paragrafo.text) < 100 and re.match(r"^\d+(?:\.\d+)?\s+", paragrafo.text):
            paragrafo.style = "Heading 2"
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragrafo.paragraph_format.first_line_indent = Inches(0)
            paragrafo.paragraph_format.space_before = Pt(18)
            paragrafo.paragraph_format.space_after = Pt(10)
        else:
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragrafo.paragraph_format.first_line_indent = Inches(0.25)
            paragrafo.paragraph_format.space_after = Pt(6)

    stile_normale = documento.styles["Normal"]
    stile_normale.font.name = "Georgia"
    stile_normale.font.size = Pt(11)
    aggiungi_numeri_pagina_docx(documento)
    output = BytesIO()
    documento.save(output)
    output.seek(0)
    return output


def estrai_anteprima_manoscritto(file_caricato) -> str:
    """Legge un estratto di DOCX/PDF senza modificare il file originale."""
    dati = BytesIO(file_caricato.getvalue())
    if file_caricato.name.lower().endswith(".docx"):
        documento = Document(dati)
        return "\n".join(paragrafo.text for paragrafo in documento.paragraphs[:100])
    # Import ritardato: la formattazione Word e le immagini restano
    # utilizzabili anche se un ambiente locale sta installando PyPDF2.
    import PyPDF2

    lettore = PyPDF2.PdfReader(dati)
    return "\n".join((pagina.extract_text() or "") for pagina in lettore.pages[:15])
